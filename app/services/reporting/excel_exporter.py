import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy.orm import Session
from app.models.exam import Exam
from app.models.session import ExamSession, StudentAnswer


def generate_excel_gradebook(db: Session, exam_id: int) -> io.BytesIO:
    """Generates official Excel spreadsheet (.xlsx) with student grades, subscores, and integrity."""
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    sessions = (
        db.query(ExamSession)
        .filter(ExamSession.exam_id == exam_id)
        .order_by(ExamSession.total_score.desc())
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekapitulasi Nilai"

    # Title styling
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    sub_font = Font(name="Calibri", size=11, italic=True)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")

    ws["A1"] = f"REKAPITULASI NILAI RESMI UJIAN: {exam.title.upper() if exam else 'UJIAN'}"
    ws["A1"].font = title_font
    ws["A2"] = f"Kelas: {exam.class_.name if exam and exam.class_ else '-'} | Guru: {exam.teacher.full_name if exam and exam.teacher else '-'}"
    ws["A2"].font = sub_font

    headers = [
        "No",
        "NISN / ID",
        "Nama Siswa",
        "Status Ujian",
        "Skor PG",
        "Skor Esai",
        "Nilai Akhir (Total)",
        "Indeks Integritas (%)",
    ]

    # Write headers at row 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Populate rows
    for row_idx, s in enumerate(sessions, start=5):
        mcq_pts = sum([
            float(a.final_score or 0.0)
            for a in s.answers
            if a.question and a.question.question_type == "multiple_choice"
        ])
        essay_pts = sum([
            float(a.final_score or 0.0)
            for a in s.answers
            if a.question and a.question.question_type != "multiple_choice"
        ])

        row_data = [
            row_idx - 4,
            s.student.identifier if s.student else "-",
            s.student.full_name if s.student else "Siswa",
            s.status.capitalize(),
            mcq_pts,
            essay_pts,
            float(s.total_score or 0.0),
            s.integrity_score,
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx in [1, 2, 4, 8]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in [5, 6, 7]:
                cell.alignment = Alignment(horizontal="right")

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
